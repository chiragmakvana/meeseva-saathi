"""Build the recordable pitch pack from the rendered slides and the narration file:

  SvaKritAI_FinalPitchDeck.pptx   16:9 deck, one full-bleed slide image per slide, the script in the speaker notes
                                  (open in PowerPoint, Slide Show > Presenter View or Record, read the notes)
  SvaKritAI_FinalPitchDeck.pdf    the same slides as a PDF (for the submission portal)
  SvaKritAI_TeamPitch_Script.pdf  printable script: one page per slide with a thumbnail, speaker, timing and the words

Usage: python scripts/build_pitch_pack.py <slides-png dir> <out dir>
The slides-png dir comes from scripts/render_slides.mjs (one <slide-id>.png per slide plus order.json).
"""
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from pitch_script import ORDER, SPEAKER, TITLE, ON_SCREEN  # noqa: E402  (same folder)

ROOT = Path(__file__).resolve().parent.parent
PNG = Path(sys.argv[1])
OUT = Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)

segs = json.loads((ROOT / "docs/demo/pitch_segments.json").read_text(encoding="utf-8"))
dur = json.loads((ROOT / "docs/demo/pitch_audio/durations.json").read_text(encoding="utf-8"))
by_slide = {}
for s in segs:
    if s["slide"] != "note":
        by_slide.setdefault(s["slide"], []).append(s)


def who(sid):
    return " then ".join(dict.fromkeys(SPEAKER[s["voice"]] for s in by_slide[sid]))


def secs(sid):
    return sum(dur[s["id"]] for s in by_slide[sid])


# ---- PPTX ----------------------------------------------------------------------------------------------------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
for n, sid in enumerate(ORDER, 1):
    sl = prs.slides.add_slide(blank)
    sl.shapes.add_picture(str(PNG / f"{sid}.png"), 0, 0, prs.slide_width, prs.slide_height)
    tf = sl.notes_slide.notes_text_frame
    tf.text = f"Slide {n} of {len(ORDER)} · {TITLE[sid]} · {who(sid)} · about {secs(sid):.0f} s"
    for s in by_slide[sid]:
        p = tf.add_paragraph(); p.text = ""
        p = tf.add_paragraph(); p.text = f"{SPEAKER[s['voice']]}:"; p.runs[0].font.bold = True
        p = tf.add_paragraph(); p.text = s["text"]; p.runs[0].font.size = Pt(14)
prs.save(OUT / "SvaKritAI_FinalPitchDeck.pptx")

# ---- PDF of the slides -----------------------------------------------------------------------------------------
imgs = [Image.open(PNG / f"{sid}.png").convert("RGB") for sid in ORDER]
imgs[0].save(OUT / "SvaKritAI_FinalPitchDeck.pdf", save_all=True, append_images=imgs[1:], resolution=150)

# ---- Printable script -------------------------------------------------------------------------------------------
W, H = A4
c = canvas.Canvas(str(OUT / "SvaKritAI_TeamPitch_Script.pdf"), pagesize=A4)
c.setTitle("MeeSeva Saathi · team pitch script")
total = sum(secs(s) for s in ORDER)


def wrap(text, font, size, width):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if c.stringWidth(t, font, size) <= width:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur:
        lines.append(cur)
    return lines


def para(x, y, text, font="Helvetica", size=11, lead=None, width=W - 40 * mm):
    lead = lead or size * 1.35
    c.setFont(font, size)
    for ln in wrap(text, font, size, width):
        c.drawString(x, y, ln); y -= lead
    return y


# cover / instructions page
y = H - 30 * mm
c.setFont("Helvetica-Bold", 22); c.drawString(20 * mm, y, "MeeSeva Saathi · team pitch script"); y -= 10 * mm
c.setFont("Helvetica", 12); c.drawString(20 * mm, y, f"Chirag and Bansi · {len(ORDER)} slides · {total / 60:.1f} min of speech · aim to finish under 8 minutes"); y -= 12 * mm
for line in [
    "How to record",
    "1. Open SvaKritAI_FinalPitchDeck.pptx in PowerPoint. Slide Show > Record (or Presenter View on a second screen): the notes pane shows these same words.",
    "2. Camera on both of you if possible; otherwise record the slides with voice only. Quiet room, phone on silent, one take per slide is fine: cut on the slide changes.",
    "3. Each page of this document is one slide: speaker, time budget, what is on screen, and the words. Read naturally; the words are a guide, not a test.",
    "4. Live-demo option at slide 5: switch to the running app and say the slide 5 and 6 words while doing runbook steps 1-4 and 9-11 (LIVE_DEMO_STEPS.md). Same 60 seconds; back to the deck at slide 7.",
    "5. Export: File > Export > Create a Video (Full HD, use recorded timings and narrations). Name it SvaKritAI_TeamPitch.mp4. Export the slides to PDF as well (already provided as SvaKritAI_FinalPitchDeck.pdf).",
    "",
    "Overview",
]:
    if line in ("How to record", "Overview"):
        c.setFont("Helvetica-Bold", 13); c.drawString(20 * mm, y, line); y -= 7 * mm
    elif line:
        y = para(20 * mm, y, line, size=10.5); y -= 2 * mm
    else:
        y -= 3 * mm
c.setFont("Helvetica", 10.5)
run = 0.0
for n, sid in enumerate(ORDER, 1):
    c.drawString(20 * mm, y, f"{n:>2}.  {TITLE[sid]}")
    c.drawRightString(W - 60 * mm, y, who(sid)); c.drawRightString(W - 38 * mm, y, f"{secs(sid):.0f} s"); c.drawRightString(W - 20 * mm, y, f"{int(run // 60)}:{int(run % 60):02d}")
    run += secs(sid); y -= 5.5 * mm
c.showPage()

# one page per slide
run = 0.0
for n, sid in enumerate(ORDER, 1):
    img = PNG / f"{sid}.png"
    iw, ih = Image.open(img).size
    tw = W - 40 * mm; th = tw * ih / iw
    c.drawImage(str(img), 20 * mm, H - 20 * mm - th, tw, th)
    y = H - 28 * mm - th
    c.setFont("Helvetica-Bold", 15); c.drawString(20 * mm, y, f"Slide {n} of {len(ORDER)} · {TITLE[sid]}"); y -= 8 * mm
    c.setFont("Helvetica", 11); c.drawString(20 * mm, y, f"{who(sid)} · about {secs(sid):.0f} s · starts near {int(run // 60)}:{int(run % 60):02d}"); y -= 7 * mm
    y = para(20 * mm, y, "On screen: " + ON_SCREEN[sid], font="Helvetica-Oblique", size=10); y -= 5 * mm
    for s in by_slide[sid]:
        c.setFont("Helvetica-Bold", 13); c.drawString(20 * mm, y, SPEAKER[s["voice"]] + ":"); y -= 7.5 * mm
        y = para(20 * mm, y, s["text"], size=14, lead=20); y -= 5 * mm
    if sid == "screens-ask":
        y = para(20 * mm, y, "Live-demo option: say these words (and slide 6's) over the running app instead of the screenshots; runbook steps 1-4 and 9-11.", font="Helvetica-Oblique", size=10)
    c.setFont("Helvetica", 9); c.drawRightString(W - 20 * mm, 12 * mm, f"{n} / {len(ORDER)}")
    run += secs(sid)
    c.showPage()
c.save()
print("pack written to", OUT, "->", [p.name for p in OUT.iterdir()])
